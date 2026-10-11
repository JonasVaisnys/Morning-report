#!/usr/bin/env python3
"""
Laikinas šaltinių patikros scriptas: tikrina, kurie naftos, dujų, CO2,
elektros kainų ir naujienų šaltiniai pasiekiami iš GitHub Actions.
"""

import io
import json
import sys
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta, timezone

import requests

UA_BROWSER = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")
UA_BOT = "Mozilla/5.0 (compatible; EnerconsultBot/1.0; +https://enerconsult.lt)"

ENERGY_KW = ["oil", "crude", "brent", "wti", "opec", "gas", "lng", "ttf",
             "power", "electricity", "energy", "carbon", "emission", "coal",
             "refin", "diesel", "hormuz", "pipeline"]

now = datetime.now(timezone.utc)
d0 = (now - timedelta(days=1)).strftime("%Y-%m-%d")
start = (now - timedelta(days=2)).strftime("%Y-%m-%dT21:00:00.000Z")
end = now.strftime("%Y-%m-%dT20:59:59.999Z")

PRICE = [
    ("Yahoo BZ=F q1", "https://query1.finance.yahoo.com/v8/finance/chart/BZ=F?range=1mo&interval=1d", "yahoo"),
    ("Yahoo BZ=F q2", "https://query2.finance.yahoo.com/v8/finance/chart/BZ=F?range=1mo&interval=1d", "yahoo"),
    ("Yahoo TTF=F", "https://query1.finance.yahoo.com/v8/finance/chart/TTF=F?range=1mo&interval=1d", "yahoo"),
    ("Yahoo EURUSD=X", "https://query1.finance.yahoo.com/v8/finance/chart/EURUSD=X?range=1mo&interval=1d", "yahoo"),
    ("Yahoo NG=F", "https://query1.finance.yahoo.com/v8/finance/chart/NG=F?range=5d&interval=1d", "yahoo"),
    ("Yahoo JKM=F", "https://query1.finance.yahoo.com/v8/finance/chart/JKM=F?range=5d&interval=1d", "yahoo"),
    ("Yahoo search carbon", "https://query2.finance.yahoo.com/v1/finance/search?q=carbon%20emissions%20futures&quotesCount=10&newsCount=0", "json"),
    ("Yahoo search EUA", "https://query2.finance.yahoo.com/v1/finance/search?q=EU%20carbon%20allowance&quotesCount=10&newsCount=0", "json"),
    ("Stooq cb.f", "https://stooq.com/q/d/l/?s=cb.f&i=d", "text"),
    ("ECB USD/EUR", "https://data-api.ecb.europa.eu/service/data/EXR/D.USD.EUR.SP00.A?lastNObservations=5&format=csvdata", "text"),
    ("Elering LT", f"https://dashboard.elering.ee/api/nps/price?start={start}&end={end}", "elering"),
    ("Nord Pool LT", f"https://dataportal-api.nordpoolgroup.com/api/DayAheadPrices?date={d0}&market=DayAhead&deliveryArea=LT&currency=EUR", "json"),
    ("Energy-charts LT", "https://api.energy-charts.info/price?bzn=LT", "json"),
    ("AGSI EU (be rakto)", "https://agsi.gie.eu/api?type=eu", "json"),
    ("EEX EUA aukcionai 2026", "https://public.eex-group.com/eex/eua-auction-report/emission-spot-primary-market-auction-report-2026-data.xlsx", "xlsx"),
    ("EIA Brent xls", "https://www.eia.gov/dnav/pet/hist_xls/RBRTEd.xls", "bytes"),
]

NEWS = [
    ("Bloomberg markets", "https://feeds.bloomberg.com/markets/news.rss"),
    ("Bloomberg green", "https://feeds.bloomberg.com/green/news.rss"),
    ("Bloomberg economics", "https://feeds.bloomberg.com/economics/news.rss"),
    ("Bloomberg energy", "https://feeds.bloomberg.com/energy/news.rss"),
    ("Bloomberg industries", "https://feeds.bloomberg.com/industries/news.rss"),
    ("Bloomberg politics", "https://feeds.bloomberg.com/politics/news.rss"),
    ("GNews oil", "https://news.google.com/rss/search?q=Brent+OR+%22crude+oil%22+OR+OPEC+when:1d&hl=en-US&gl=US&ceid=US:en"),
    ("GNews gas", "https://news.google.com/rss/search?q=%22European+gas%22+OR+TTF+OR+LNG+when:1d&hl=en-US&gl=US&ceid=US:en"),
    ("GNews bloomberg energy", "https://news.google.com/rss/search?q=site:bloomberg.com+(oil+OR+gas+OR+LNG)+when:1d&hl=en-US&gl=US&ceid=US:en"),
    ("OilPrice", "https://oilprice.com/rss/main"),
    ("FT energy", "https://www.ft.com/energy?format=rss"),
    ("CNBC energy", "https://www.cnbc.com/id/19836768/device/rss/rss.html"),
    ("S&P oil", "https://www.spglobal.com/commodityinsights/en/rss-feed/oil"),
    ("S&P gas", "https://www.spglobal.com/commodityinsights/en/rss-feed/natural-gas"),
    ("S&P LNG", "https://www.spglobal.com/commodityinsights/en/rss-feed/lng"),
    ("Euractiv energy", "https://www.euractiv.com/sections/energy-environment/feed/"),
    ("LNG Prime", "https://lngprime.com/feed/"),
    ("Rigzone", "https://www.rigzone.com/news/rss/rigzone_latest.aspx"),
    ("EIA Today in Energy", "https://www.eia.gov/rss/todayinenergy.xml"),
    ("EIA press", "https://www.eia.gov/rss/press_rss.xml"),
]


def get(url, ua=UA_BROWSER):
    return requests.get(url, headers={"User-Agent": ua, "Accept": "*/*"}, timeout=20)


def show_head(r):
    ct = r.headers.get("content-type", "?")
    print(f"    HTTP {r.status_code} · {ct} · {len(r.content)} B")
    if r.status_code != 200:
        print(f"    pradžia: {r.text[:160]!r}")


def probe_price(name, url, kind):
    print(f"\n### {name}\n    {url}")
    try:
        r = get(url)
    except Exception as exc:
        print(f"    KLAIDA: {exc}")
        return
    show_head(r)
    if r.status_code != 200:
        return
    try:
        if kind == "yahoo":
            res = r.json()["chart"]["result"][0]
            m = res["meta"]
            print(f"    meta: {m.get('symbol')} {m.get('currency')} {m.get('exchangeName')} "
                  f"price={m.get('regularMarketPrice')} time={m.get('regularMarketTime')}")
            ts = res.get("timestamp", [])
            closes = res["indicators"]["quote"][0].get("close", [])
            for t, c in list(zip(ts, closes))[-6:]:
                print(f"      {datetime.fromtimestamp(t, timezone.utc):%Y-%m-%d %H:%M} {c}")
        elif kind == "elering":
            j = r.json()
            data = j.get("data", {})
            print(f"    raktai: {list(j.keys())} · zonos: {list(data.keys())}")
            lt = data.get("lt", [])
            print(f"    lt įrašų: {len(lt)}; pirmi: {lt[:2]}; paskutiniai: {lt[-2:]}")
        elif kind == "json":
            j = r.json()
            if isinstance(j, dict):
                print(f"    raktai: {list(j.keys())[:15]}")
                if "quotes" in j:
                    for q in j["quotes"][:10]:
                        print(f"      {q.get('symbol')} · {q.get('shortname')} · {q.get('exchange')} · {q.get('quoteType')}")
                if "multiAreaEntries" in j:
                    e = j["multiAreaEntries"]
                    print(f"    įrašų: {len(e)}; pirmas: {e[:1]}")
                    print(f"    areaAverages: {j.get('areaAverages')}")
                if "price" in j and "unix_seconds" in j:
                    print(f"    energy-charts: {len(j['price'])} reikšmių; paskutinės: {j['price'][-4:]}")
            else:
                print(f"    sąrašas: {len(j)}; {str(j)[:200]}")
        elif kind == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
            print(f"    lapai: {wb.sheetnames}")
            ws = wb[wb.sheetnames[0]]
            rows = list(ws.iter_rows(values_only=True))
            print(f"    eilučių: {len(rows)}")
            for row in rows[:8]:
                print(f"      {row[:12]}")
            print("    ...")
            for row in rows[-3:]:
                print(f"      {row[:12]}")
        elif kind == "text":
            lines = r.text.strip().splitlines()
            print(f"    eilučių: {len(lines)}")
            for ln in lines[:2] + ["..."] + lines[-4:]:
                print(f"      {ln[:200]}")
        else:
            print(f"    pradžia: {r.content[:80]!r}")
    except Exception as exc:
        print(f"    PARSAVIMO KLAIDA: {exc}; pradžia: {r.text[:200]!r}")


def probe_news(name, url):
    print(f"\n### {name}\n    {url}")
    for ua_name, ua in (("bot", UA_BOT), ("browser", UA_BROWSER)):
        try:
            r = get(url, ua)
        except Exception as exc:
            print(f"    [{ua_name}] KLAIDA: {exc}")
            continue
        print(f"    [{ua_name}]", end="")
        show_head(r)
        if r.status_code == 200:
            break
    else:
        return
    try:
        root = ET.fromstring(r.content)
    except ET.ParseError as exc:
        print(f"    XML KLAIDA: {exc}; pradžia: {r.text[:200]!r}")
        return
    ns = {"atom": "http://www.w3.org/2005/Atom"}
    items = root.findall(".//item") or root.findall(".//atom:entry", ns)
    n_energy = 0
    for it in items:
        title = (it.findtext("title") or it.findtext("atom:title", namespaces=ns) or "")
        if any(k in title.lower() for k in ENERGY_KW):
            n_energy += 1
    print(f"    elementų: {len(items)} · energetikos pagal pavadinimą: {n_energy}")
    for it in items[:6]:
        title = (it.findtext("title") or it.findtext("atom:title", namespaces=ns) or "").strip()
        pub = (it.findtext("pubDate") or it.findtext("atom:updated", namespaces=ns) or "").strip()
        src = it.find("source")
        src_txt = f" · šaltinis: {src.text} ({src.get('url')})" if src is not None else ""
        desc = (it.findtext("description") or "").strip()
        print(f"      - {title[:110]} | {pub[:31]}{src_txt} | aprašas {len(desc)} simb.")


if __name__ == "__main__":
    print(f"Patikra {now:%Y-%m-%d %H:%M} UTC")
    print("\n================ KAINOS ================")
    for p in PRICE:
        probe_price(*p)
    print("\n================ NAUJIENOS ================")
    for n in NEWS:
        probe_news(*n)
